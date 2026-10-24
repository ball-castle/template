#!/usr/bin/env python3
"""Personal contest notes: source-backed code, model cards and real page references."""
import argparse
from datetime import date
import math
import re
import unicodedata
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm
from build_notebooks import setup_doc, font
from round02_common import ROOT, CATALOG, as_text, cards, catalogs, code_digest
from round02_common import digest, expanded_source, json_write, load_reports, read_json, verified

OUTPUT = ROOT/'板子打印版'
PAGE_BUDGET = 725  # A4 body is about 768 pt; retain measured physical slack.
FIELDS = [('signal','识别'),('observe','先看什么'),('brute','从暴力出发'),
          ('redundancy','重复与性质'),('derive','推出模型'),('conditions','适用条件'),
          ('proof','一句证明'),('dp_state','状态含义'),('transition','转移'),
          ('initial','初值'),('order','遍历顺序'),('complexity','复杂度'),
          ('pitfalls','反例与易错'),('exercise','练习与来源')]
CONTAINERS = [
    ('每次取当前最小 / 最大','小根堆 / 大根堆'),
    ('动态有序去重 / 允许重复','set / multiset'),
    ('动态频率计数','map / unordered_map'),
    ('查前驱、后继、第一个不小于x','set / multiset'),
    ('先进先出 / 后进先出','queue / stack'),
    ('两端增删 / 固定窗口候选','deque / 单调队列'),
    ('字符串前缀 / 最大异或数对','字符串Trie / 二进制Trie'),
    ('单点加、前缀与区间和','Fenwick'),
    ('静态区间min/max/gcd','ST表'),
    ('动态区间合并信息 / 整段修改','线段树 / Lazy线段树'),
    ('连通关系 / 无权最短路','DSU / BFS'),
    ('0/1权 / 非负一般权最短路','0-1 BFS / Dijkstra'),
]

def width(text):
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in text)

def clean(value):
    s = as_text(value)
    s = re.sub(r'https://github.com/ball-castle/(cf|template)/blob/[^/]+/',r'\1/',s)
    return s.replace('https://codeforces.com/problemset/problem/','CF ')

def bookmark(paragraph, name, number):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end = OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    paragraph._p.insert(0,start);paragraph._p.append(end)

class Writer:
    def __init__(self,title):
        self.doc=setup_doc(title)
        self.doc.sections[0].header.paragraphs[0].clear()
        font(self.doc.sections[0].header.paragraphs[0].add_run(
            'ball-castle 个人竞赛手册  |  GNU++20  |  '+date.today().isoformat()),size=8)
        self.doc.styles['Code'].paragraph_format.line_spacing=Pt(10.2)
        self.pages=0;self.used=0;self.section='';self.code=[];self.records=[]
        self.bookmark_id=1

    def page(self,title,record=None,continued=False):
        needs_break=bool(self.pages)
        self.pages+=1;self.used=32;self.section=title
        p=self.doc.add_paragraph(style='Heading 1')
        p.paragraph_format.page_break_before=needs_break
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(7)
        font(p.add_run(title+(' 续' if continued else '')),size=14,bold=True)
        if record is not None:
            marker='R2_'+re.sub(r'\W','_',record['id'])
            bookmark(p,marker,self.bookmark_id);self.bookmark_id+=1
            record['bookmark']=marker;record['planned_page']=self.pages
            self.records.append(record)

    def para(self,text,bold=False,size=9.5):
        text=clean(text)
        cost=13*sum(max(1,math.ceil(width(line)/103)) for line in text.split('\n'))+4
        if self.used+cost>PAGE_BUDGET:
            self.page(self.section,continued=True)
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=Pt(13)
        font(p.add_run(text),size=size,bold=bold);self.used+=cost

    def compact_section(self,title,record,estimated):
        if self.used+estimated+42>PAGE_BUDGET:
            self.page(title,record);return
        p=self.doc.add_paragraph(style='Heading 1')
        p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(7)
        font(p.add_run(title),size=14,bold=True)
        marker='R2_'+re.sub(r'\W','_',record['id'])
        bookmark(p,marker,self.bookmark_id);self.bookmark_id+=1
        record['bookmark']=marker;record['planned_page']=self.pages
        self.records.append(record);self.section=title;self.used+=42

    def code_lines(self,lines):
        for i,line in enumerate(lines):
            if getattr(self,'patch',False) and re.match(r'^(?:inline\s+)?[A-Za-z_][A-Za-z_0-9:<>, ]*\s+\w+\s*\(',line):
                balance=0;opened=False;end=None
                for j in range(i,min(len(lines),i+20)):
                    if '{' in lines[j]:opened=True
                    balance+=lines[j].count('{')-lines[j].count('}')
                    if opened and balance==0:end=j+1;break
                if end is not None:
                    block=sum(10.2*max(1,math.ceil(width(x)/99)) for x in lines[i:end])
                    if self.used+block>PAGE_BUDGET and self.used>32:
                        self.page(self.section,continued=True)
            cost=10.2*max(1,math.ceil(width(line)/99))
            # Keep a short final block together within the measured page height.
            tail_cost=sum(10.2*max(1,math.ceil(width(x)/99)) for x in lines[i:]) if len(lines)-i<=8 else 10000
            budget=max(PAGE_BUDGET,self.used+tail_cost) if len(lines)-i<=5 else (745 if self.used+tail_cost<=745 else PAGE_BUDGET)
            if self.used+cost>budget:
                self.page(self.section,continued=True)
            p=self.doc.add_paragraph(style='Code')
            p.paragraph_format.line_spacing=Pt(10.2)
            p.paragraph_format.keep_with_next=(line.rstrip().endswith('(') or
                (i+1<len(lines) and lines[i+1].strip() in ('}','};')))
            font(p.add_run(line),family='Consolas',size=9)
            self.code.append(line);self.used+=cost

    def table(self,rows,widths):
        table=self.doc.add_table(rows=0,cols=len(widths));table.autofit=False
        for row in rows:
            cells=table.add_row().cells
            for cell,content,colwidth in zip(cells,row,widths):
                cell.width=Mm(colwidth)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=Pt(12)
                font(p.add_run(str(content)),size=9)
        self.used+=len(rows)*19
        return table

def selection(kind,old,new):
    chosen_cards=cards()
    all_modules=old|new
    if kind=='patch':
        config=read_json(CATALOG/'patch_selection.json')
        wanted=set(config['cards'])
        chosen_cards=[c for c in chosen_cards if c['id'] in wanted]
        missing=wanted-{c['id'] for c in chosen_cards}
        if missing:raise RuntimeError('Missing patch cards '+str(missing))
        module_ids=config['modules']
    else:
        replaced={x for m in new.values() for x in m.get('replaces',[])}
        module_ids=[k for k in old if k not in replaced]+list(new)
    chosen_modules=[all_modules[x] for x in module_ids]
    # Same real file has one printed source; keep every teaching ID as an alias.
    groups={}
    for m in chosen_modules:
        key=m['source']
        if key not in groups:groups[key]={'module':m,'aliases':[]}
        groups[key]['aliases'].append(m['id'])
    return chosen_cards,list(groups.values()),all_modules

def build(kind):
    old,new=catalogs();chosen_cards,groups,all_modules=selection(kind,old,new)
    reports=load_reports()
    for group in groups:
        if not verified(group['module'],reports):
            raise RuntimeError('No current verified source/test binding: '+group['module']['id'])
    title='成都站 赛前小技巧与题型补丁 2026' if kind=='patch' else 'ACM ICPC 铜牌稳妥覆盖 算法与模型 V2'
    stem='成都站_赛前小技巧与题型补丁_2026' if kind=='patch' else 'ACM_ICPC_铜牌稳妥覆盖_算法与模型_V2'
    w=Writer(title);w.patch=(kind=='patch');w.page(title)
    w.para('先从题意找模型，再查代码。优先直接采用我原来验证使用过的实现，只修具体问题；基础树存储和容器用法与算法同样重要。',bold=True)
    w.para('这是短补充资料，配合第一轮24页急用版使用；未在本册展开的源码按编号到仓库查阅。' if kind=='patch' else
           '本册按识别信号讲推导，再给个人比赛代码和调用；全部模型、来源与测试可在仓库按同一编号检索。')
    w.para('原版inline ll read()保留；按题目n/T读取常规整数，其他输入用cin。关闭iostream同步后，不混用cin与getchar。')
    w.para('纸质资料可携带数量、规格：待赛站确认。生成文件不表示全部允许带入比赛。',bold=True)
    w.para('先按操作选容器',bold=True,size=11)
    w.table([('题目需要','首先考虑')]+CONTAINERS,[98,82])
    w.para('复杂度先看n、q、值域和状态数；数量级只是筛选，不能替代证明或实测。')
    w.page('从题意到实现的六步')
    for head,body in [
      ('1 看约束','很小先穷举；n≈40考虑折半；集合小考虑状压；大n寻找线性或nlogn；值域小计数、大值域离散化。'),
      ('2 看操作','连续段→前缀/差分/数据结构；最少步→下界构造/BFS/DP；可达→不变量+充分性；轮流行动→先核对终止与胜负规则。'),
      ('3 写暴力','枚举合法对象，找重复计算；问未来究竟还需要哪些历史信息。'),
      ('4 补证明','构造要合法且达到目标；贪心写交换；二分证单调；DP证状态足够并划分互斥完整。'),
      ('5 选源码','先用已有熟悉代码。写清1/0下标、闭/半开区间、初始化、值域、输出含义。'),
      ('6 找反例','最小规模、全相等、全负、边界、断图/重边、多测；小数据暴力与优化对拍，记录最小反例。')]:
        w.para(head,bold=True,size=11);w.para(body)
    w.para('DP状态检查：保存什么／漏了什么／初值和不可达／转移是否重计／遍历依赖／答案位置／max、min还是sum／是否取模／能否滚动。',bold=True)
    w.para('WA先核对语义与边界；RE查下标、空容器和递归深度；TLE先重算复杂度；多测清实际使用过的状态。')
    w.para('每个代码模块单独使用。同名全局a/g/tree/init等来自不同模板，拼到同一题前只保留所需实现或按题目改名。本地#include在打印中已递归展开；无需照抄文件依赖。')
    w.page('我的代码习惯与 DP 状态识别')
    w.para('原代码验证过就直接采用：ll、1下标、全局a/g/dp/cnt、独立函数；原Fenwick的struct/vector同样保留。树用vector<ll> g[N]存双向边；新题先init清上一组状态。',bold=True)
    w.para('线段树沿用pushup、pushdown、build、modify、query；DSU沿用init、find、unite、same。已验证的原实现与新增修复均有来源记录，不为统一外观改写。')
    w.table([('识别信号','必须保存的状态 / 编号'),
      ('一段只能接前一段','以i结尾的最优值 DPC01'),('只能右下走','到(i,j)的方案数/代价 DPC02'),
      ('每件最多一次 / 无限 / 有限 / 分组','容量与阶段，遍历方向不同 DPC03–06'),
      ('最后一次合并把区间切开','区间[l,r]与切分点k DPC07'),
      ('父子选择互斥','子树根选/不选 DPC08'),('依赖无环','拓扑序上的最优前驱 DPC09'),
      ('小集合与当前位置','mask和last DPC10'),('区间很大但数位少','pos/tight/started和题目状态 DPC11'),
      ('保留原序 / 两序列匹配','末尾值或两个前缀长度 DPC12'),
      ('历史只影响当前阶段','位置与有限状态 DPC13'),
      ('问方案数 / 随机收益','互斥计数 / 概率加权与自环移项 DPC14–15')],[90,90])
    w.para('状态相同必须拥有相同的后续选择。max/min保留最优；计数相加但防重；不可达要单独表示。详推导与完整代码见V2相同编号。')
    if kind=='full':
        w.page('XOR / AND / OR 思维地图')
        w.para('XOR是逐位模2加：同数两次抵消。连续区间先写前缀XOR；求等于k的数量查freq[prefix^k]，先查后加，起点freq[0]=1。')
        w.para('数对贡献：某位c个1、n-c个0，XOR有c(n-c)对，AND有C(c,2)对，OR有C(n,2)-C(n-c,2)对。计数值、移位位宽和累加溢出要分别检查。')
        w.para('最大数对/子数组→二进制Trie；最大子集XOR→线性基；value[u]^value[v]=w→图上相对值；独立无环正常规则游戏→SG异或。')
        w.para('构造先列长度、值域、互异、正数等全部约束；最后一个数由目标XOR倒推，只保证XOR式，不保证其他限制。')
        w.para('从原题提取的思维缺口：2268B补全可达性；2268C目标实际是区间XOR AND区间最大值；2218E先核对sum n<=3105，原双循环可用；1457D证明压缩到小n；909F分开证明两种排列；2261B给popcount不变量补充分性。详细个人复盘见XOR08、12–15与MY_XOR_REVIEW.md。')
    # Index fields are filled by Word after bookmarks exist, then checked against actual PDF.
    index_records=[('C_'+c['id'],c['id']+' '+c['title']) for c in chosen_cards]
    index_records += [('M_'+g['module']['id'],' / '.join(g['aliases'])+' '+g['module']['title']) for g in groups]
    for start in range(0,len(index_records),30):
        w.page('题型识别与代码索引'+(' 续' if start else ''))
        rows=[('稳定编号与内容','页')]+[(label,'') for _,label in index_records[start:start+30]]
        table=w.table(rows,[165,15])
        for row,(key,_) in zip(table.rows[1:],index_records[start:start+30]):
            p=row.cells[1].paragraphs[0]
            field=OxmlElement('w:fldSimple')
            field.set(qn('w:instr'),'PAGEREF R2_'+re.sub(r'\W','_',key)+' \\h')
            run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='?'
            run.append(text);field.append(run);p._p.append(field)
    for card in chosen_cards:
        record={'id':'C_'+card['id'],'kind':'card','catalog_id':card['id'],'title':card['title']}
        fields=FIELDS if kind=='full' else [('signal','识别'),('derive','怎么推出'),('proof','关键证明'),('pitfalls','别套错'),('exercise','练习')]
        heading='【'+card['id']+'】'+card['title']
        if kind=='patch':
            estimate=sum(13*max(1,math.ceil(width(clean(card.get(k,''))+label)/103))+4 for k,label in fields)
            estimate+=13*max(1,math.ceil(width(clean(card.get('origins',[])))/103))+30
            usage=card.get('usage','')
            if usage:estimate+=13*sum(max(1,math.ceil(width(x)/99)) for x in clean(usage).split('\n'))
            # Start the first model on a new page after the index.
            if not w.records:w.page(heading,record)
            else:w.compact_section(heading,record,estimate)
        else:w.page(heading,record)
        for key,label in fields:
            if card.get(key):w.para(label+'：'+clean(card[key]))
        if kind=='full' and card.get('personal_extension'):
            w.para('我的原题延伸：'+clean(card['personal_extension']))
        refs=card.get('code_refs',[])
        unresolved=[x for x in refs if x not in all_modules]
        if unresolved:raise RuntimeError(card['id']+' unknown references '+str(unresolved))
        w.para('代码入口：'+('、'.join(refs) if refs else '思维框架，结合本卡具体练习；不是可直接粘贴的算法。'),bold=True)
        usage=card.get('usage')
        if usage:
            if isinstance(usage,str) and re.search(r'\bint\s+main\s*\(',usage):
                for report in reports:
                    if report.get('card_examples',{}).get(card['id'])!=code_digest(usage.splitlines()):
                        raise RuntimeError('Unverified runnable card example '+card['id'])
                w.code_lines(usage.splitlines())
                record['example_code']=usage.splitlines()
            else:w.para('调用提醒：'+clean(usage))
        if card.get('origins'):w.para('原代码：'+clean(card['origins']),size=8)
    for group in groups:
        module=group['module'];lines,sources=expanded_source(module['source'])
        record={'id':'M_'+module['id'],'kind':'module','catalog_id':module['id'],
                'aliases':group['aliases'],'title':module['title'],'source':module['source'],
                'source_hashes':sources,'expanded_code_sha256':code_digest(lines),'code_lines':lines}
        w.page('【'+' / '.join(group['aliases'])+'】'+module['title'],record)
        for key,label in [('when','识别'),('idea','思路'),('proof','证明'),('contract','约定'),
                          ('usage','调用'),('example','例子'),('complexity','复杂度'),('pitfalls','易错')]:
            if module.get(key):
                if kind=='patch' and key in ('idea','example','pitfalls'):
                    continue
                if module['id']=='IO01' and key in ('contract','pitfalls'):
                    continue
                w.para(label+'：'+clean(module[key]))
        origins=module.get('origins',[])
        if origins:w.para('来源：'+'；'.join(o['path'] for o in origins),size=8)
        w.para('规范源码：'+module['source']+'；完整独立代码如下；测试：'+clean(module['test']),size=8)
        w.code_lines(lines)
        if module.get('example_source'):
            example_path=module['example_source']
            for report in reports:
                evidence=report.get('module_examples',{}).get(module['id'],{})
                if evidence.get('sha256')!=digest(ROOT/example_path,text=True):
                    raise RuntimeError('Unverified module example '+module['id'])
            example_lines=[line for line in (ROOT/example_path).read_text(encoding='utf-8').splitlines()
                if not re.fullmatch(r'\s*#\s*include\s*"[^"]+"\s*',line)]
            w.para('接在上述代码后面的完整main（已独立运行样例）：',bold=True)
            w.para('样例输入：'+module.get('sample_input','').replace('\n',' / ')+
                   '\n样例输出：'+module.get('sample_output','').replace('\n',' / '))
            w.code_lines(example_lines)
            record['example_source']=example_path
            record['example_sha256']=digest(ROOT/example_path,text=True)
            record['example_code']=example_lines
    if kind=='full':
        # API-specific teaching is part of the paper book, not merely an online coverage name.
        w.page('基础 STL 逐项速查：场景、调用、复杂度与坑')
        w.para('以下每项都有受测调用；完整代码按模块编号查前面的代码索引。几个字符就能使用的API在此集中比较，避免为每个API另建算法文件。')
        for topic in read_json(CATALOG/'stl_coverage.json'):
            if topic['group'][0] not in ('A','B'):continue
            values=[('用途',topic['scenario']),('实际调用',topic['actual_calls']),
                    ('复杂度',topic['complexity']),('易错',topic['pitfalls']),
                    ('已检查例子',topic['validation']['evidence'])]
            record={'id':'T_'+topic['id'],'catalog_id':'TIP_'+topic['id'],
                    'kind':'tip','title':topic['topic'],'coverage_id':topic['id'],'indexed':False}
            estimate=sum(13*max(1,math.ceil(width(clean(v)+k)/103))+4 for k,v in values)+35
            w.compact_section('【TIP_'+topic['id']+'】'+topic['topic'],record,estimate)
            for label,value in values:w.para(label+'：'+clean(value))
            w.para('代码与来源：'+'、'.join(topic['code_refs'])+'；'+
                '、'.join(x['source'] for x in topic.get('code',[]) if x.get('source')),size=8)
    OUTPUT.mkdir(exist_ok=True)
    target=OUTPUT/(stem+'.docx');w.doc.save(target)
    manifest={'book':stem,'kind':kind,'planned_pages':w.pages,
              'code_lines':len(w.code),'expected_code_lines':w.code,
              'sections':w.records,'word_pdf_binding_required':True,
              'source_hash_mode':'LF-normalized','date':date.today().isoformat()}
    json_write(OUTPUT/(stem+'.round02.json'),manifest)
    print(f'{target}: {w.pages} planned pages, {len(w.records)} sections, {len(w.code)} exact code lines')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--book',choices=['patch','full'],required=True)
    build(parser.parse_args().book)
