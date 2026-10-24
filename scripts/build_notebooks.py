#!/usr/bin/env python3
"""Build editable DOCX from verified sources and structured metadata."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / '板子打印版'
FONT = 'Microsoft YaHei'

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def text(value):
    if isinstance(value, list):
        return '；'.join(map(str, value))
    if isinstance(value, dict):
        return '；'.join(f'{k}: {v}' for k, v in value.items())
    return str(value)

def clean(value):
    # Keep long URLs in the searchable digital index, use human identifiers on paper.
    return re.sub(r'https?://\S+', lambda m: m[0].replace('https://github.com/ball-castle/cf/blob/main/', 'cf/')
                  .replace('https://codeforces.com/problemset/problem/', 'CF ')
                  .replace('https://www.luogu.com.cn/problem/', '洛谷 '), text(value))

def font(run, family=FONT, size=9.5, bold=False):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)

def paragraph(doc, content, bold=False, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(13)
    font(p.add_run(content), size=size, bold=bold)
    return p

def heading(doc, title, first=False):
    p = doc.add_paragraph(style='Title' if first else 'Heading 1')
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    font(p.add_run(title), size=20 if first else 14, bold=True)

def new_page(doc, title):
    doc.add_page_break()
    heading(doc, title)

def code_line(doc, line):
    p = doc.add_paragraph(style='Code')
    font(p.add_run(line), family='Consolas', size=9.25)

def setup_doc(name):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(13), Mm(13)
    sec.left_margin, sec.right_margin = Mm(13), Mm(13)
    sec.header_distance, sec.footer_distance = Mm(5), Mm(5)
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = Pt(13)
    normal.paragraph_format.space_after = Pt(4)
    for style_name in ['Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = FONT
    # Some runtime starter documents carry a colored Title paragraph border.
    # Remove it at the style source rather than relying on Word theme overrides.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    style = doc.styles.add_style('Code', 1)
    style.font.name, style.font.size = 'Consolas', Pt(9.25)
    style.paragraph_format.line_spacing = Pt(10.8)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.widow_control = False
    style.paragraph_format.keep_together = True
    font(sec.header.paragraphs[0].add_run(name + '  |  GNU++20  |  2026-10-24'), size=8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('第 '), size=8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    font(footer.add_run(' 页'), size=8)
    doc.core_properties.title = name
    doc.core_properties.subject = 'CF 1500 选手的 ICPC 铜牌复习与经过验证的算法源码'
    doc.core_properties.author = 'ball-castle'
    return doc

def source_chunks(module):
    lines = (ROOT / module['source']).read_text(encoding='utf-8-sig').splitlines()
    # First page reserves room for the usage contract. Continued pages have larger capacity.
    first = 51 if len(lines) == 51 else 50
    if module['id'] == 'CO02':
        first = next(i for i, line in enumerate(lines) if line.startswith('// 离散化'))
    split_before = {'DS03':'    ll query(', 'G05':'    Unicycle r;',
                    'G06':'inline DijkstraResult dijkstra(', 'G09':'    int distance('}
    if module['id'] in split_before:
        first = next(i for i, line in enumerate(lines) if line.startswith(split_before[module['id']]))
    chunks = [lines[:first]]
    rest = lines[first:]
    while rest:
        chunks.append(rest[:65])
        rest = rest[65:]
    return chunks

def add_index(doc, modules, pages, page_num):
    paragraph(doc, '先识别题型、写出证明，再按稳定编号查模板。所有正式代码都由仓库源码自动生成。')
    paragraph(doc, '赛站纸质资料数量与规格：待赛站确认。25 页不是已核实的成都站限制。', bold=True)
    paragraph(doc, '统一提醒：多测新建容器；类型与区间以各模块契约为准；所有算术结果须处于声明范围。')
    paragraph(doc, '思维卡片见第 2–3 页；排错与最后复习见第 4 页。', bold=True)
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Mm(19), Mm(146), Mm(15)]
    for cell, label, width in zip(table.rows[0].cells, ['编号', '算法与识别入口', '页码'], widths):
        cell.width = width
        font(cell.paragraphs[0].add_run(label), size=9, bold=True)
    for module in modules:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, [module['id'], module['title'], str(pages[module['id']])], widths):
            cell.width = width
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = Pt(11)
            font(p.add_run(str(value)), size=8.8)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side)
                for key, value in {'val':'single','sz':'4','color':'D9D9D9'}.items():
                    node.set(qn('w:'+key), value)
                borders.append(node)
            props.append(borders)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EEEEEE')
                props.append(shade)

def build(kind, all_modules, reports):
    modules = [m for m in all_modules if kind == 'full' or m['priority'] == 'P0']
    for module in modules:
        for report in reports:
            if report['status'] != 'PASS' or report['sources'].get(module['source']) != digest(ROOT/module['source']):
                raise RuntimeError('Unverified or changed source: '+module['source'])
            tests = re.findall(r'tests/[A-Za-z0-9_]+\.cpp', text(module['test']))
            if not tests:
                raise RuntimeError('Missing test declaration: '+module['id'])
            for test in tests:
                if report.get('test_sources', {}).get(test) != digest(ROOT/test):
                    raise RuntimeError('Test driver changed or was not run: '+test)
    stem = '成都站_铜牌冲刺板子_2026' if kind == 'emergency' else 'ACM_ICPC_铜牌算法板子_完整版'
    title = '成都站 铜牌冲刺板子 2026' if kind == 'emergency' else 'ACM ICPC 铜牌算法板子 完整版'
    # Full index is two pages when needed, with front cards retaining predictable locations.
    pages, current = {}, 5
    if len(modules) > 36:
        current += 1
    for module in modules:
        pages[module['id']] = current
        current += len(source_chunks(module))
    doc = setup_doc(title)
    heading(doc, title, first=True)
    if kind == 'full':
        paragraph(doc, '手敲时按说明一并抄入依赖模块；合并进同一 .cpp 后省略本地头文件 #include 与 #pragma once，保留 bits/stdc++.h。cp:: 按调用示例使用。', size=9)
    add_index(doc, modules[:36], pages, 1)
    front = json.loads((ROOT/'notes/print_front.json').read_text(encoding='utf-8'))
    for item in front:
        new_page(doc, item['title'])
        for section_title, body, example in item['sections']:
            paragraph(doc, section_title, bold=True, size=11)
            paragraph(doc, body)
            paragraph(doc, example, size=9)
            paragraph(doc, '', size=3)
    if len(modules) > 36:
        new_page(doc, '算法索引 续')
        add_index(doc, modules[36:], pages, 5)
    printed_lines = []
    for module in modules:
        chunks = source_chunks(module)
        for chunk_index, lines in enumerate(chunks):
            suffix = f' 续 {chunk_index+1}' if chunk_index else ''
            new_page(doc, f"{module['id']} {module['title']}{suffix}")
            if not chunk_index:
                paragraph(doc, '识别：'+clean(module['when']))
                paragraph(doc, '约定：'+clean(module['contract']))
                paragraph(doc, '复杂度：'+clean(module['complexity']))
                paragraph(doc, '调用：'+clean(module['usage']))
                paragraph(doc, '易错：'+clean(module['pitfalls']))
                paragraph(doc, '例题：'+clean(module['example']), size=8.5)
                paragraph(doc, '源码 '+module['source']+' | 测试 '+text(module['test'])+' | 独立编译及所属套件通过；ASan/UBSan 已运行', size=8)
            else:
                paragraph(doc, '接上一页；请合并复制整个模块。源码 '+module['source'], size=8)
            for line in lines:
                code_line(doc, line)
                printed_lines.append(line)
    OUTPUT.mkdir(exist_ok=True)
    path = OUTPUT / (stem+'.docx')
    doc.save(path)
    # A byte-for-byte line comparison catches omissions and independent print-code forks.
    with ZipFile(path) as z:
        xml = ET.fromstring(z.read('word/document.xml'))
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    actual = []
    for p in xml.findall('.//w:body/w:p', ns):
        style = p.find('w:pPr/w:pStyle', ns)
        if style is not None and style.get(qn('w:val')) == 'Code':
            actual.append(''.join(p.itertext()))
    if actual != printed_lines:
        raise RuntimeError('DOCX source parity failure')
    manifest = {'book':stem, 'planned_pages': current-1, 'module_count':len(modules),
                'code_lines':len(printed_lines), 'code_parity':'PASS',
                'modules':[{'id':m['id'],'source':m['source'],'sha256':digest(ROOT/m['source']),
                            'start_page':pages[m['id']]} for m in modules]}
    (OUTPUT/(stem+'.manifest.json')).write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'{path}: planned {current-1} pages, {len(modules)} modules, exact source parity PASS')

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--book', choices=['emergency','full','both'], default='both')
    args = ap.parse_args()
    modules = [json.loads(p.read_text(encoding='utf-8-sig')) for p in sorted((ROOT/'notes/modules').glob('*.json'))]
    modules.sort(key=lambda m: (['P0','P1','P2'].index(m['priority']), m['id']))
    if not modules:
        raise RuntimeError('No verified modules')
    reports = [json.loads((ROOT/'build'/mode/'results.json').read_text(encoding='utf-8'))
               for mode in ['release','sanitizers']]
    for kind in ['emergency','full'] if args.book == 'both' else [args.book]:
        build(kind, modules, reports)

if __name__ == '__main__':
    main()
